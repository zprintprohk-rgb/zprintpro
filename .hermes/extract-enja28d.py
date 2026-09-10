#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""提取 en/ja 28d 热门查询 (imp>=3) → .hermes/enja28d-queries.json
用途: 标题 v4 补满方案 3 筛选 (GSC 实证) 校验"""
import openpyxl, os, json, sys
sys.stdout.reconfigure(encoding='utf-8')
DIR = r'F:\zprintpro-nextjs\GSC数据'
def clean_key(k):
    k = str(k)
    for src, dst in [('点击率','ctr'),('点击次数','clicks'),('热门查询','query'),('搜索查询','query'),
                     ('排名靠前的网页','page'),('着陆页面','page'),('点击','clicks'),('展示','imp'),
                     ('排名','pos'),('网页','page'),('页面','page'),('查询','query')]:
        if src in k: return dst
    return k
def load_queries(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    qs = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        if not rows: continue
        header = [clean_key(h) for h in rows[0]]
        if 'query' not in header: continue
        for r in rows[1:]:
            if not any(c is not None for c in r): continue
            d = dict(zip(header, r))
            q = d.get('query')
            if not q or not str(q).strip(): continue
            qs.append({'query': str(q), 'imp': int(d.get('imp') or 0), 'clicks': int(d.get('clicks') or 0), 'pos': round(float(d.get('pos') or 0), 1)})
    return qs
en = load_queries(os.path.join(DIR, 'zprintpro.com-Performance-on-Search-28天美国站点数据2026-09-10.xlsx'))
ja = load_queries(os.path.join(DIR, 'zprintpro.com-Performance-on-Search-28天日本站点数据2026-09-10.xlsx'))
# 去重 + 按 imp 排序
def dedup(qs):
    seen = {}
    for q in qs:
        k = q['query'].lower()
        if k not in seen or q['imp'] > seen[k]['imp']: seen[k] = q
    return sorted(seen.values(), key=lambda x: -x['imp'])
en = dedup(en); ja = dedup(ja)
json.dump({'en': [q for q in en if q['imp'] >= 3], 'ja': [q for q in ja if q['imp'] >= 3]}, open(r'F:\zprintpro-nextjs\.hermes\enja28d-queries.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('=== en 28d 热门 (imp>=10) ===')
for q in [q for q in en if q['imp'] >= 10][:40]:
    print(f"  {q['imp']:>4} imp {q['pos']:>6} pos | {q['query']}")
print('=== ja 28d 热门 (imp>=10) ===')
for q in [q for q in ja if q['imp'] >= 10][:40]:
    print(f"  {q['imp']:>4} imp {q['pos']:>6} pos | {q['query']}")
