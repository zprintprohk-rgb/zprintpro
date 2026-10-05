# -*- coding: utf-8 -*-
"""10/5 vs 9/29 同窗口 pos 变化 + 10/5 头部词全景
输出 .hermes/gsc-2026-10-05-week-delta.json
"""
import io, json
from openpyxl import load_workbook

BASE = r'F:\zprintpro-nextjs\GSC数据'
NEW = {
    '7d-sum':  BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
    '28d-sum': BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
    'hk-7d':   BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
    'us-7d':   BASE + r'\美国站点7天数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
    'jp-7d':   BASE + r'\日本站点7天数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
}
OLD = {
    '7d-sum':  BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    '28d-sum': BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-7d':   BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'us-7d':   BASE + r'\美国站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'jp-7d':   BASE + r'\日本站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}

def load(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        q = str(r[0]).strip()
        try:
            out[q] = {'clicks': int(float(r[1] or 0)), 'imps': int(float(r[2] or 0)),
                      'pos': round(float(r[4]), 2) if r[4] is not None else None}
        except Exception:
            continue
    return out

result = {'windows': {}, 'deltas': {}}
for name in NEW:
    new, old = load(NEW[name]), load(OLD[name])
    top = dict(sorted(new.items(), key=lambda kv: -kv[1]['imps'])[:70])
    result['windows'][name] = {'total_q': len(new), 'top': top}
    deltas = []
    for q, v in new.items():
        o = old.get(q)
        if o and o.get('pos') and v.get('pos') and v['imps'] >= 2:
            d = round(o['pos'] - v['pos'], 1)
            if abs(d) >= 3:
                deltas.append({'q': q, 'old_pos': o['pos'], 'new_pos': v['pos'], 'delta': d,
                               'imps': v['imps'], 'clicks': v['clicks']})
    result['deltas'][name] = sorted(deltas, key=lambda x: -x['imps'])[:40]

io.open(r'F:\zprintpro-nextjs\.hermes\gsc-2026-10-05-week-delta.json', 'w', encoding='utf-8').write(
    json.dumps(result, ensure_ascii=False, indent=1))
for name in NEW:
    w = result['windows'][name]
    print(f"== {name}: {w['total_q']} queries, top5: " + ', '.join(
        f"{q} {v['imps']}@{v['pos']}" for q, v in list(w['top'].items())[:5]))
    for d in result['deltas'][name][:10]:
        print(f"   {'+' if d['delta']>0 else ''}{d['delta']}  {d['q']}  {d['old_pos']}->{d['new_pos']}  imps={d['imps']}")
