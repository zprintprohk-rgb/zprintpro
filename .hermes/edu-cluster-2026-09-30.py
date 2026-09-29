# -*- coding: utf-8 -*-
"""教育簇头词裁决: 24h/7d/28d/3mo 四窗口 GSC 全量提取
输出: .hermes/gsc-2026-09-30-edu-cluster.json
"""
import io, json, re
from openpyxl import load_workbook

BASE = r'F:\zprintpro-nextjs\GSC数据'
FILES = {
    '24h':  BASE + r'\24小时三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    '7d':   BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    '28d':  BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    '3mo':  BASE + r'\3个月三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}
SITE_FILES = {
    'hk-7d':  BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-28d': BASE + r'\香港站点28天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-3mo': BASE + r'\香港站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    'us-3mo': BASE + r'\美国站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'jp-3mo': BASE + r'\日本站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}

EDU_RX = re.compile(r'教育|校園|學校|学校|校簿|校薄|證書|証書|证书|証明書|証明|卒業|成績|school|education|certificate|workbook|textbook|handbook|教科書|ワークブック|diploma|transcript|入学|入学|招生', re.I)

def load_queries(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for r in rows[1:]:
        q = r[0]
        if not q:
            continue
        q = str(q).strip()
        try:
            clicks = float(r[1] or 0); imps = float(r[2] or 0)
            ctr = float(r[3]) if r[3] is not None else None
            pos = float(r[4]) if r[4] is not None else None
        except Exception:
            continue
        out[q] = {'clicks': int(clicks), 'imps': int(imps), 'ctr': ctr, 'pos': round(pos, 2) if pos else None}
    return out

result = {}
for name, path in {**FILES, **SITE_FILES}.items():
    d = load_queries(path)
    cluster = {q: v for q, v in d.items() if EDU_RX.search(q)}
    result[name] = {
        'total_queries': len(d),
        'edu_queries': len(cluster),
        'edu_imps': sum(v['imps'] for v in cluster.values()),
        'queries': dict(sorted(cluster.items(), key=lambda kv: -kv[1]['imps']))
    }

io.open(r'F:\zprintpro-nextjs\.hermes\gsc-2026-09-30-edu-cluster.json', 'w', encoding='utf-8').write(
    json.dumps(result, ensure_ascii=False, indent=1))
for name in result:
    r = result[name]
    print(f"{name}: edu_q={r['edu_queries']} edu_imps={r['edu_imps']} (of {r['total_queries']} total q)")
