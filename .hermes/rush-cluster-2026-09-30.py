# -*- coding: utf-8 -*-
"""即日急件/ rush 簇提取: 即日|急件|特急|当日|same day|rush|24h 等
输出: .hermes/gsc-2026-09-30-rush-cluster.json
"""
import io, json, re
from openpyxl import load_workbook

BASE = r'F:\zprintpro-nextjs\GSC数据'
FILES = {
    '24h':  BASE + r'\24小时三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    '7d':   BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    '28d':  BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    '3mo':  BASE + r'\3个月三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}
SITE_FILES = {
    'hk-7d':  BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-28d': BASE + r'\香港站点28天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-3mo': BASE + r'\香港站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    'us-3mo': BASE + r'\美国站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'jp-3mo': BASE + r'\日本站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}

RUSH_RX = re.compile(
    r'即日|急件|特急|当日|即場|即印|速印|快速印刷|加急'
    r'|same[ -]?day|next[ -]?day|rush|24[ -]?h|24hour|overnight|urgent|fast print|quick print|express print'
    r'|当日納品|即納|即日印刷|特急印刷|24時間|翌日',
    re.I)

def load_queries(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for r in rows[1:]:
        q = r[0]
        if not q:
            continue
        q = str(q).strip()
        try:
            clicks = float(r[1] or 0); imps = float(r[2] or 0)
            ctr = float(r[3]) if r[3] is not None else None
            pos = float(r[4]) if r[4] is not None else None
        except Exception:
            continue
        out[q] = {'clicks': int(clicks), 'imps': int(imps), 'ctr': ctr, 'pos': round(pos, 2) if pos else None}
    return out

result = {}
for name, path in {**FILES, **SITE_FILES}.items():
    d = load_queries(path)
    cluster = {q: v for q, v in d.items() if RUSH_RX.search(q)}
    result[name] = {
        'total_queries': len(d),
        'rush_q': len(cluster),
        'rush_imps': sum(v['imps'] for v in cluster.values()),
        'queries': dict(sorted(cluster.items(), key=lambda kv: -kv[1]['imps']))
    }

io.open(r'F:\zprintpro-nextjs\.hermes\gsc-2026-09-30-rush-cluster.json', 'w', encoding='utf-8').write(
    json.dumps(result, ensure_ascii=False, indent=1))
for name in result:
    r = result[name]
    print(f"{name}: rush_q={r['rush_q']} rush_imps={r['rush_imps']} (of {r['total_queries']})")
