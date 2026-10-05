# -*- coding: utf-8 -*-
"""K3 大词清单提取: 10/5 7d-sum + hk-7d + 28d-sum vs 9/29"""
import io, json, re
from openpyxl import load_workbook

BASE = r'F:\zprintpro-nextjs\GSC数据'
RX = re.compile(r'海報印刷|即日急件|印刷公司|校簿|簿印刷|簿 印刷|車身廣告|摺頁|易拉架|學校印刷|学校印刷|急件印刷|印書|貼紙訂製|食品包裝|書刊印刷|a2.*即日|即日.*a2|a2 印刷|印刷 即日|roll.?up|海報.*即日|即日.*海報|印海報', re.I)

FILES = {
    '7d-sum': (BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
               BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx'),
    'hk-7d': (BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
              BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx'),
    '28d-sum': (BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-10-05.xlsx',
                BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx'),
}

def load(p):
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r[0]:
            continue
        q = str(r[0]).strip()
        try:
            out[q] = {'imps': int(float(r[2] or 0)), 'clicks': int(float(r[1] or 0)),
                      'pos': round(float(r[4]), 2) if r[4] is not None else None}
        except Exception:
            continue
    return out

result = {}
for name, (newp, oldp) in FILES.items():
    new, old = load(newp), load(oldp)
    rows = []
    for q, v in new.items():
        if RX.search(q) and v['imps'] >= 1:
            o = old.get(q, {})
            rows.append({'q': q, 'imps': v['imps'], 'clicks': v['clicks'], 'pos': v['pos'],
                         'old_pos': o.get('pos')})
    result[name] = sorted(rows, key=lambda x: -x['imps'])

io.open(r'F:\zprintpro-nextjs\.hermes\gsc-2026-10-05-k3-bigwords.json', 'w', encoding='utf-8').write(
    json.dumps(result, ensure_ascii=False, indent=1))
for name in FILES:
    print('==', name)
    for r in result[name][:24]:
        d = ('%+.1f' % (r['old_pos'] - r['pos'])) if r.get('old_pos') and r.get('pos') else '  -  '
        print(f"  {r['pos']!s:>6} (was {r['old_pos']!s:>6} {d})  imps={r['imps']:<4} {r['q']}")
