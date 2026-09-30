# -*- coding: utf-8 -*-
"""新增作战簇提取: 小册子/可变数据/流水码/白墨 + 关联长尾
输出: .hermes/gsc-2026-09-30-newwords-cluster.json
"""
import io, json, re
from openpyxl import load_workbook

BASE = r'F:\zprintpro-nextjs\GSC数据'
FILES = {
    '24h':  BASE + r'\24小时三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    '7d':   BASE + r'\7天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    '28d':  BASE + r'\28天三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    '3mo':  BASE + r'\3个月三站点汇总数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}
SITE_FILES = {
    'hk-7d':  BASE + r'\香港站点7天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-28d': BASE + r'\香港站点28天数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'hk-3mo': BASE + r'\香港站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29 (1).xlsx',
    'us-3mo': BASE + r'\美国站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
    'jp-3mo': BASE + r'\日本站点3个月数据zprintpro.com-Performance-on-Search-2026-09-29.xlsx',
}

NEW_RX = re.compile(
    r'小册子|小冊子|冊子|书册|書冊|booklet|pamphlet'
    r'|可变数据|可變數據|可変|variable|流水|序列号|序列號|连番|連番|serial|sequential|numbering|编号|編號|ナンバリング'
    r'|白墨|white[ -]?ink|ホワイト',
    re.I)

def load_queries(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for r in rows[1:]:
        q = r[0]
        if not q:
            continue
        q = str(q).strip()
        try:
            clicks = float(r[1] or 0); imps = float(r[2] or 0)
            ctr = float(r[3]) if r[3] is not None else None
            pos = float(r[4]) if r[4] is not None else None
        except Exception:
            continue
        out[q] = {'clicks': int(clicks), 'imps': int(imps), 'ctr': ctr, 'pos': round(pos, 2) if pos else None}
    return out

result = {}
for name, path in {**FILES, **SITE_FILES}.items():
    d = load_queries(path)
    cluster = {q: v for q, v in d.items() if NEW_RX.search(q)}
    result[name] = {
        'total_queries': len(d),
        'new_q': len(cluster),
        'new_imps': sum(v['imps'] for v in cluster.values()),
        'queries': dict(sorted(cluster.items(), key=lambda kv: -kv[1]['imps']))
    }

io.open(r'F:\zprintpro-nextjs\.hermes\gsc-2026-09-30-newwords-cluster.json', 'w', encoding='utf-8').write(
    json.dumps(result, ensure_ascii=False, indent=1))
for name in result:
    r = result[name]
    print(f"{name}: new_q={r['new_q']} new_imps={r['new_imps']} (of {r['total_queries']})")
