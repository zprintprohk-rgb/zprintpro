#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
scripts/gsc-cliff-detector.py — GSC 断崖检测器 (spam 恢复防护机制, K3 2026-10-12)

背景: 2026-09 spam 更新 10/8 完成推送后整站命中 (诊断 v2, commit fcdd5b57)。
恢复期内需要每日机器判据: ①有没有新的断崖 ②恢复曲线是否在爬升 ③数据是否新鲜。

数据源: GSC数据/<站点>站点28天数据*.xlsx (sheet「图表」= 逐日 日期/点击/展示/CTR/排名)
        + 可选 实时24小时*.xlsx (逐小时, 全 property, 最鲜信号)

判据 (§0.35.4 机器判据, 不看自我播报):
  CRITICAL_CLIFF : 最新完整日展示 < 前 7 日均值 × 30%       → exit 1
  WARN_DEPRESSED : 近 3 日均值 < 28 天日均 × 50% (持续低位, 恢复期跟踪用, 不拦)
  WARN_STALE     : 最新数据日期距今天 > 5 天 (GSC 自身滞后 ~3 天 + 导出 slack 2 天)
  WARN_NO_FILE   : 某站点找不到 28 天档

用法: python scripts/gsc-cliff-detector.py [--json]   # exit 1 = 有 CRITICAL
"""
import glob
import json
import re
import sys
import warnings
from datetime import date, datetime
from pathlib import Path

warnings.filterwarnings('ignore')  # openpyxl 「无默认样式」UserWarning 无害

ROOT = Path(__file__).resolve().parent.parent
GSC_DIR = ROOT / 'GSC数据'
LOG_DIR = ROOT / '.hermes' / 'logs'
CLIFF_RATIO = 0.30
DEPRESSED_RATIO = 0.50
STALE_DAYS = 5
SITES = {'香港': 'hk', '日本': 'jp', '美国': 'us'}
# 24h 实时档命名兼容: 旧导出 = 实时24小时*; K3 10-15 起新导出 = 24小时三站点汇总*
REALTIME_PATTERNS = ('实时24小时*', '24小时三站点汇总*')
# 断崖前基线 (K3 10-13 拍板口径, 固定常量): gsc-cliff-2026-10-13.json 三站 mean28 之和
# (hk 377.2 + jp 72.2 + us 135.8 = 585.2/日, 28d 档止于 10-05 = 命中前窗口).
# 固定分母保证逐日恢复系数可比 ("连续 3 日 >=60%" 判定的前提); 动态 mean28 比值保留为次要参考.
PRECLIFF_BASELINE_COMBINED = 585.2
PRECLIFF_BASELINE_SOURCE = 'gsc-cliff-2026-10-13.json sites.{hk,jp,us}.mean28_imps (28d 档止 10-05)'


def latest_file_any(patterns):
    cands = []
    for pat in patterns:
        cands.extend(glob.glob(str(GSC_DIR / pat)))
    if not cands:
        return None
    return Path(sorted(cands, key=lambda p: _file_date(p), reverse=True)[0])


def latest_file(pattern):
    files = sorted(glob.glob(str(GSC_DIR / pattern)), key=lambda p: _file_date(p), reverse=True)
    return Path(files[0]) if files else None


def _file_date(p):
    m = re.search(r'(\d{4}-\d{2}-\d{2})', Path(p).name)
    return m.group(1) if m else '0000-00-00'


def read_daily(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb['图表'] if '图表' in wb.sheetnames else wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        if r and r[0] and re.match(r'^\d{4}-\d{2}-\d{2}$', str(r[0])):
            rows.append({'date': str(r[0]), 'clicks': float(r[1] or 0), 'imps': float(r[2] or 0),
                         'ctr': float(r[3] or 0), 'pos': float(r[4] or 0)})
    return rows


def read_hourly(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb['图表'] if '图表' in wb.sheetnames else wb.active
    rows = []
    for r in ws.iter_rows(values_only=True):
        if r and r[0] and re.match(r'^\d{4}-\d{2}-\d{2}T', str(r[0])):
            rows.append({'hour': str(r[0]), 'clicks': float(r[1] or 0), 'imps': float(r[2] or 0)})
    return rows


def main():
    today = date.today()
    report = {'date': today.isoformat(), 'sites': {}, 'realtime_24h': None, 'alerts': []}

    for site_zh, code in SITES.items():
        f = latest_file(f'{site_zh}站点28天数据*.xlsx')
        caliber = 'single-site'
        fallback = None
        if f:
            _frows = read_daily(f)
            _flage = (date.today() - datetime.strptime(_frows[-1]['date'], '%Y-%m-%d').date()).days if _frows else 999
            if _flage > STALE_DAYS:
                # A17 (2026-10-29): 单站档过旧 (age>5) 但 K3 已改导三站点汇总形态 ->
                # 回退读汇总档 (全国家口径) 避免 WARN_STALE 误报; 记录双口径证据。
                fb = latest_file('28天三站点汇总数据*.xlsx')
                if fb:
                    fbrows = read_daily(fb)
                    fb_age = (date.today() - datetime.strptime(fbrows[-1]['date'], '%Y-%m-%d').date()).days if fbrows else 999
                    if fbrows and fb_age < _flage:
                        fallback = {'old': f, 'old_age': _flage, 'new': fb, 'new_age': fb_age}
                        f = fb
                        caliber = 'full-country'
        if not f:
            # A17 修复 (2026-10-29): K3 自 10-15 后改导「三站点汇总」形态 (全国家口径),
            # 单站形态不再出新档 -> 回退读汇总档作为全站序列, 标注 caliber 防口径混算。
            f = latest_file('28天三站点汇总数据*.xlsx')
            caliber = 'full-country'
            if f:
                report['alerts'].append({'level': 'INFO_CALIBER_SWITCH', 'site': code,
                                         'detail': f'单站档缺失, 回退汇总档 {f.name} (全国家口径, 非单站)'})
        if not f:
            report['alerts'].append({'level': 'WARN_NO_FILE', 'site': code, 'detail': '找不到 28 天档'})
            continue
        if fallback:
            report['alerts'].append({'level': 'INFO_CALIBER_SWITCH', 'site': code,
                                     'detail': f'单站档 {fallback["old"].name} 过旧 {fallback["old_age"]}d, 回退汇总档 '
                                               f'{fallback["new"].name} ({fallback["new_age"]}d, 全国家口径)'})
        rows = read_daily(f)
        if len(rows) < 8:
            report['alerts'].append({'level': 'WARN_NO_FILE', 'site': code, 'detail': f'{f.name} 行数不足 ({len(rows)})'})
            continue
        last = rows[-1]
        prior7 = rows[-8:-1]
        prior7_mean = sum(r['imps'] for r in prior7) / len(prior7)
        trail3 = rows[-3:]
        trail3_mean = sum(r['imps'] for r in trail3) / len(trail3)
        mean28 = sum(r['imps'] for r in rows) / len(rows)
        last_date = datetime.strptime(last['date'], '%Y-%m-%d').date()
        age = (today - last_date).days

        s = {
            'file': f.name, 'days': len(rows), 'latest_date': last['date'], 'data_age_days': age,
            'latest_imps': last['imps'], 'prior7d_mean_imps': round(prior7_mean, 1),
            'trailing3d_mean_imps': round(trail3_mean, 1), 'mean28_imps': round(mean28, 1),
            'total28_imps': sum(r['imps'] for r in rows), 'total28_clicks': sum(r['clicks'] for r in rows),
            'cliff_ratio': round(last['imps'] / prior7_mean, 3) if prior7_mean else None,
            'depression_ratio': round(trail3_mean / mean28, 3) if mean28 else None,
            'caliber': caliber,
        }
        report['sites'][code] = s

        if age > STALE_DAYS:
            report['alerts'].append({'level': 'WARN_STALE', 'site': code,
                                     'detail': f'最新数据 {last["date"]} 已 {age} 天 (> {STALE_DAYS}), 需导出新档'})
        if prior7_mean >= 20 and last['imps'] < prior7_mean * CLIFF_RATIO:
            report['alerts'].append({'level': 'CRITICAL_CLIFF', 'site': code,
                                     'detail': f'{last["date"]} 展示 {last["imps"]:.0f} < 前7日均 {prior7_mean:.0f} × {CLIFF_RATIO} — 新断崖!'})
        if mean28 >= 20 and trail3_mean < mean28 * DEPRESSED_RATIO:
            report['alerts'].append({'level': 'WARN_DEPRESSED', 'site': code,
                                     'detail': f'近3日均 {trail3_mean:.0f} = 28天日均 {mean28:.0f} 的 {trail3_mean / mean28 * 100:.0f}% (持续低位)'})

    rt = latest_file_any(REALTIME_PATTERNS)
    if rt:
        hours = read_hourly(rt)
        if hours:
            rt_imps = sum(h['imps'] for h in hours)
            # A17 修复 (2026-10-29): 单站档缺失时 sites 的 mean28 已切换为全国家口径,
            # 与固定 585.2 (三站和, 单站口径) 基线不可比 -> 动态分母改用同口径全国家 mean28 并标注。
            calibers = {s['caliber'] for s in report['sites'].values()}
            base_is_full_country = bool(calibers) and calibers == {'full-country'}
            if base_is_full_country:
                # 三键读同一全站序列: 求和会把分母放大 3 倍 (749.2×3=2247.6 伪值),
                # 正确分母 = 单序列 mean28 (749.2) -> ratio 与车道 109.9%/107.2% 同量级。
                dynamic_denom = next(iter(report['sites'].values()))['mean28_imps']
            else:
                dynamic_denom = sum(s['mean28_imps'] for s in report['sites'].values())
            dynamic_denom = dynamic_denom if dynamic_denom else None
            dynamic_denom_caliber = 'full-country' if base_is_full_country else 'single-site-sum'
            ratio_dynamic = round(rt_imps / dynamic_denom, 3) if dynamic_denom else None
            # 固定断崖前基线仍按原口径输出, 但加 caliber 标签防与车道 109.9% (全国家) 混算
            fixed_ratio = round(rt_imps / PRECLIFF_BASELINE_COMBINED, 3) if PRECLIFF_BASELINE_COMBINED else None
            report['realtime_24h'] = {
                'file': rt.name, 'hours': len(hours),
                'window': f'{hours[0]["hour"]} ~ {hours[-1]["hour"]}',
                'imps': rt_imps, 'clicks': sum(h['clicks'] for h in hours),
                # 主指标: 固定断崖前基线分母 (逐日可比, "连续 3 日" 判定依据)
                'recovery_ratio_vs_precliff': fixed_ratio,
                'recovery_ratio_vs_precliff_caliber': 'single-site-sum (hk+jp+us 三站和 585.2, 与全国家 714.71 不可比)',
                # 次要参考: 动态 28d mean28 分母 (窗口随导出滑动, 日间不可比)
                'recovery_ratio_vs_mean28_dynamic': ratio_dynamic,
                'recovery_ratio_vs_mean28_dynamic_caliber': dynamic_denom_caliber,
                'baseline': {'combined_mean28': PRECLIFF_BASELINE_COMBINED, 'source': PRECLIFF_BASELINE_SOURCE},
                'caliber': 'full-country' if base_is_full_country else 'single-site',
            }

    LOG_DIR.mkdir(parents=True, exist_ok=True)
    out = LOG_DIR / f'gsc-cliff-{today.isoformat()}.json'
    out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')

    critical = [a for a in report['alerts'] if a['level'] == 'CRITICAL_CLIFF']
    for code, s in report['sites'].items():
        print(f"[cliff] {code}: 最新日 {s['latest_date']} 展示 {s['latest_imps']:.0f} | 前7日均 {s['prior7d_mean_imps']:.0f} | "
              f"近3日均 {s['trailing3d_mean_imps']:.0f} ({(s['depression_ratio'] or 0) * 100:.0f}% of 28d日均) | 数据龄 {s['data_age_days']}d")
    if report['realtime_24h']:
        r = report['realtime_24h']
        rr = r.get('recovery_ratio_vs_precliff')
        print(f"[cliff] realtime24h ({r['window']}): 展示 {r['imps']:.0f} / 点击 {r['clicks']:.0f}"
              + (f" | 恢复系数 {rr * 100:.1f}% (vs 断崖前日均, 目标爬回 100%)" if rr is not None else ""))
    for a in report['alerts']:
        print(f"[cliff] {a['level']} [{a.get('site', '-')}] {a['detail']}")
    print(f"[cliff] 报告 -> {out}")

    if '--json' in sys.argv:
        print(json.dumps(report, ensure_ascii=False))
    sys.exit(1 if critical else 0)


if __name__ == '__main__':
    main()
