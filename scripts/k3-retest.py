# -*- coding: utf-8 -*-
"""11-06 七/十四/二十八天复测器 — 对 10-30 KPI 基线出逐项 delta
用法:
  python scripts/k3-retest.py --xlsx "GSC数据/<新导出>.xlsx" [--baseline .hermes/logs/k3-kpi-panel-2026-10-30.json]
自检: 传 10-30 同一档 → 所有 delta 应为 0
"""
import sys, io, json, argparse, re, os
from openpyxl import load_workbook
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

TARGETS = {
    "海报类目": "/zh-hk/category/posters/",
    "贴纸PDP": "/en/product/small-batch-stickers/",
    "catalog服务页": "/en/services/catalog-printing-china/",
    "C109活动页": "/ja/comiket-c109/",
    "包装类目": "/zh-hk/category/packaging/",
}
MONEY_STRICT = {
    "zh-hk": ["海報印刷", "印海報", "食品包裝印刷", "貼紙印刷", "月曆印刷", "小冊子印刷", "紙袋印刷", "騎馬釘印刷", "即日印刷"],
    "en": ["small batch sticker printing", "small batch stickers", "doujinshi printing", "saddle stitch booklet printing", "china catalog printing", "large envelopes", "a2 poster", "corporate holiday cards"],
    "ja": ["コミケ 印刷", "卒アル販売", "教材 印刷製本", "教材 印刷 製本", "pvc シール", "年賀状印刷"],
}

def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True) if not all(v is None for v in r)]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--baseline", default=".hermes/logs/k3-kpi-panel-2026-10-30.json")
    a = ap.parse_args()

    base = json.load(open(a.baseline, encoding="utf-8"))
    wb = load_workbook(a.xlsx, data_only=True)

    ch = rows(wb["图表"])[1:]
    co = rows(wb["国家_地区"])[1:]
    pg = rows(wb["网页"])[1:]
    qq = rows(wb["查询数"])[1:]

    def agg(rs):
        c = sum(r[1] or 0 for r in rs); i = sum(r[2] or 0 for r in rs)
        pw = sum((r[4] or 0) * (r[2] or 0) for r in rs)
        return c, i, (c / i * 100 if i else 0), (pw / i if i else 0)

    A, B = agg(ch), agg(co)
    print(f"# 复测报告 — {os.path.basename(a.xlsx)}")
    print(f"窗口: {ch[0][0]} → {ch[-1][0]}  |  基线: {base.get('window','?')}\n")

    bl = base["four_metrics"]
    print("## 四指标 delta")
    print("| 指标 | 基线 | 本次 | delta |")
    print("|---|---|---|---|")
    print(f"| 点击 | {bl['clicks']:.0f} | {A[0]:.0f} | {A[0]-bl['clicks']:+.0f} |")
    print(f"| 展示 | {bl['impressions']:,.0f} | {A[1]:,.0f} | {A[1]-bl['impressions']:+,.0f} |")
    print(f"| CTR | {bl['ctr']:.2f}% | {A[2]:.2f}% | {A[2]-bl['ctr']:+.2f}pp |")
    print(f"| 均位 | {bl['avg_pos']:.2f} | {A[3]:.2f} | {A[3]-bl['avg_pos']:+.2f} |")
    print(f"\n双方法: 图表 {A[0]:.0f}/{A[1]:,.0f} vs 国家 {B[0]:.0f}/{B[1]:,.0f} → {'一致' if (A[0],A[1])==(B[0],B[1]) else '不一致(先查指标)'}\n")

    print("## 任务落地页（网页 sheet）")
    print("| 页面 | 点击 | 展示 | CTR | 均位 | 基线点击 | 基线展示 |")
    print("|---|---|---|---|---|---|---|")
    base_pages = {}
    for p in base.get("pages_top", []) or []:
        base_pages[p["url"]] = p
    for name, frag in TARGETS.items():
        rec = next((r for r in pg if frag in str(r[0])), None)
        if rec:
            c, i, ctr, pos = rec[1] or 0, rec[2] or 0, (rec[3] or 0) * 100, rec[4] or 0
            bp = base_pages.get(rec[0], {})
            print(f"| {name} | {c:.0f} | {i:,.0f} | {ctr:.2f}% | {pos:.2f} | {bp.get('clk','—')} | {bp.get('imp','—')} |")
        else:
            print(f"| {name} | 无数据(未进 top 页) | — | — | — | — | — |")

    print("\n## 带钱词（严格口径）")
    bs = {m["word"]: m["hit"] for m in base["money_words"]["strict"]}
    p1 = 0; total = 0
    print("| locale | 词 | 基线 pos | 本次 pos | delta | 本次 clk/imp |")
    print("|---|---|---|---|---|---|")
    for loc, words in MONEY_STRICT.items():
        for w in words:
            total += 1
            wn = re.sub(r"\s+", " ", w.strip().lower())
            rec = next((r for r in qq if re.sub(r"\s+", " ", str(r[0]).strip().lower()) == wn), None)
            cur = (round(rec[4], 2) if rec and rec[4] else None)
            if cur and cur <= 10: p1 += 1
            b = bs.get(w) or {}
            bpos = b.get("pos") if b else None
            d = f"{cur - bpos:+.2f}" if (cur and bpos) else "—"
            ci = f"{rec[1] or 0:.0f}/{rec[2] or 0:.0f}" if rec else "—"
            print(f"| {loc} | {w} | {bpos if bpos else '—'} | {cur if cur else '—'} | {d} | {ci} |")
    print(f"\n首页词数（严格）: 基线 {base['money_words']['on_page1_strict']}/{total} → 本次 {p1}/{total}")

if __name__ == "__main__":
    main()
